"""
OCR Engine
==========
Tesseract OCR ve PyMuPDF görüntü işleme köprüsü.
Taranmış PDF'lerden metin çıkarma ve Word (DOCX) belgesi üretme yeteneği sunar.
Framework-agnostic (Qt bağımlılığı yoktur).
"""

from __future__ import annotations

import io
import os
import shutil
from pathlib import Path
from typing import List, Optional


_COMMON_TESSERACT_PATHS = [
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    os.path.expanduser(r"~\AppData\Local\Programs\Tesseract-OCR\tesseract.exe"),
    os.path.expanduser(r"~\AppData\Local\Tesseract-OCR\tesseract.exe"),
    r"C:\tools\tesseract\tesseract.exe",
]


class OcrEngine:
    """Tesseract OCR işlemlerini yöneten motor."""

    def __init__(self, preferred_lang: str = "tur+eng"):
        self._preferred_lang = preferred_lang
        self._tesseract_cmd: Optional[str] = self._find_tesseract_binary()
        self._available_cache: Optional[bool] = None

    # ------------------------------------------------------------------ #
    #  Durum & Keşif                                                      #
    # ------------------------------------------------------------------ #

    def is_available(self) -> bool:
        """
        Sistemde Tesseract binary ve pytesseract'in hazır olup olmadığını kontrol eder.
        Sonuç instance ömrü boyunca cache'lenir — `pytesseract.get_tesseract_version()`
        her çağrıda bir `tesseract.exe --version` subprocess'i başlatıyor; bu property
        UI tarafından converter seçimi gibi sık tetiklenen olaylarda tekrar tekrar
        okunuyor, cache'siz hali gözle görülür donmaya yol açıyordu.
        """
        if self._available_cache is not None:
            return self._available_cache

        if not self._tesseract_cmd:
            self._available_cache = False
            return False
        try:
            import pytesseract
            from PIL import Image  # noqa: F401
            pytesseract.pytesseract.tesseract_cmd = self._tesseract_cmd
            # Basit bir sürüm sorgusu ile çalışabilirliği doğrula
            pytesseract.get_tesseract_version()
            self._available_cache = True
        except Exception:
            self._available_cache = False
        return self._available_cache

    @property
    def tesseract_path(self) -> Optional[str]:
        return self._tesseract_cmd

    def _find_tesseract_binary(self) -> Optional[str]:
        which_cmd = shutil.which("tesseract") or shutil.which("tesseract.exe")
        if which_cmd and os.path.exists(which_cmd):
            return which_cmd

        for p in _COMMON_TESSERACT_PATHS:
            if os.path.exists(p):
                return p

        return None

    # ------------------------------------------------------------------ #
    #  OCR Metin Çıkarımı                                                 #
    # ------------------------------------------------------------------ #

    def ocr_page(self, page, dpi: int = 300) -> str:
        """PyMuPDF (fitz) sayfasını yüksek çözünürlüklü işleyip OCR metnini döner."""
        import fitz
        import pytesseract
        from PIL import Image

        pytesseract.pytesseract.tesseract_cmd = self._tesseract_cmd

        # Yüksek kaliteli pixmap render
        pix = page.get_pixmap(dpi=dpi)
        img_bytes = pix.tobytes("png")
        img = Image.open(io.BytesIO(img_bytes))

        # Önce tercih edilen dili (tur+eng) dene; dil paketi yoksa varsayılana düş
        try:
            text = pytesseract.image_to_string(img, lang=self._preferred_lang)
        except Exception:
            try:
                text = pytesseract.image_to_string(img, lang="eng")
            except Exception:
                text = pytesseract.image_to_string(img)

        return text.strip()

    def ocr_pdf_to_text(self, source_path: Path, dpi: int = 300) -> str:
        """Tüm PDF sayfalarını OCR'dan geçirip birleştirilmiş metin döner."""
        import fitz

        doc = fitz.open(str(source_path))
        page_texts: List[str] = []
        try:
            for i, page in enumerate(doc):
                txt = self.ocr_page(page, dpi=dpi)
                page_texts.append(f"--- Sayfa {i + 1} ---\n\n{txt}")
            return "\n\n".join(page_texts)
        finally:
            doc.close()

    # ------------------------------------------------------------------ #
    #  OCR → DOCX Belge Üretimi                                          #
    # ------------------------------------------------------------------ #

    def ocr_pdf_to_docx(
        self,
        source_path: Path,
        output_path: Path,
        dpi: int = 250,
    ) -> int:
        """
        Taranmış PDF sayfalarını OCR ile okuyup düzenlenebilir temiz bir
        Word (DOCX) belgesi oluşturur.
        """
        import fitz
        from docx import Document
        from docx.shared import Pt, Inches

        doc = fitz.open(str(source_path))
        word_doc = Document()

        # Sayfa kenar boşlukları
        for section in word_doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        total_pages = len(doc)
        try:
            for i, page in enumerate(doc):
                text = self.ocr_page(page, dpi=dpi)

                if i > 0:
                    word_doc.add_page_break()

                # Sayfa başlığı
                p_hdr = word_doc.add_paragraph()
                r_hdr = p_hdr.add_run(f"Sayfa {i + 1}")
                r_hdr.font.size = Pt(8.5)
                r_hdr.font.italic = True
                r_hdr.font.color.rgb = None

                # Metin paragraflarını ekle
                lines = [line.strip() for line in text.split("\n") if line.strip()]
                for line in lines:
                    p = word_doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.line_spacing = 1.15
                    run = p.add_run(line)
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)

            output_path.parent.mkdir(parents=True, exist_ok=True)
            word_doc.save(str(output_path))
            return total_pages
        finally:
            doc.close()
