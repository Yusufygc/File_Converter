"""
DOCX → TXT Converter
=======================
`python-docx` (zaten bir bağımlılık — bkz. requirements.txt) ile
paragraf metnini çıkarır. LibreOffice'e ihtiyaç duymaz — saf Python,
tek satırlık metin çıkarma işlemi için harici süreç başlatmak gereksiz.

Tablo/başlık/altbilgi kapsanmaz (yalnızca `document.paragraphs`) —
bilinçli bir sınırlama, ilk sürüm için yeterli.

SRP : Yalnızca DOCX → TXT metin çıkarmadan sorumlu.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional

from core.converters.base import BaseConverter, ConvertOutcome
from core.interfaces.converter_interface import ConversionOptions


class DocxToTxtConverter(BaseConverter):
    """DOCX → TXT dönüştürücü. Yalnızca paragraf metnini çıkarır."""

    # ------------------------------------------------------------------ #
    #  IConverter interface                                                #
    # ------------------------------------------------------------------ #

    @property
    def source_extension(self) -> str:
        return ".docx"

    @property
    def target_extension(self) -> str:
        return ".txt"

    @property
    def display_name(self) -> str:
        return "Word → TXT"

    def _do_convert(
        self,
        source_path: Path,
        output_path: Path,
        options: ConversionOptions,
    ) -> Optional[ConvertOutcome]:
        from docx import Document

        doc = Document(str(source_path))
        text = "\n".join(p.text for p in doc.paragraphs)
        output_path.write_text(text, encoding="utf-8")

        return None

    @property
    def unavailable_hint(self) -> str:
        return (
            "Dönüşüm motoru bulunamadı.\n\n"
            "python-docx (önerilen):\n"
            "  pip install python-docx"
        )

    # ------------------------------------------------------------------ #
    #  Public helpers                                                      #
    # ------------------------------------------------------------------ #

    @property
    def is_available(self) -> bool:
        try:
            import docx  # noqa: F401
            return True
        except ImportError:
            return False

    @property
    def active_engine_name(self) -> str:
        return "python-docx" if self.is_available else "Yok"

    @property
    def is_parallel_safe(self) -> bool:
        # Saf Python, dış süreç/paylaşımlı durum yok.
        return True
